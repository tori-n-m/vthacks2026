from io import BytesIO
import json
from pathlib import Path

from flask import Flask, jsonify, render_template, request, send_file
from docx import Document

from ai.ai_service import list_styles, summarize_document, tutor_chat

app = Flask(__name__)


@app.get("/styles")
def styles():
    return jsonify(list_styles())


@app.post("/summarize")
def summarize():
    uploaded = request.files.get("document")
    style = request.form.get("style", "general")
    describe_images = request.form.get("describe_images") == "on"
    socratic_tutor = request.form.get("socratic_tutor") == "on"

    if uploaded is None or not uploaded.filename:
        return jsonify({"ok": False, "error": "Choose a document to upload."}), 400

    result = summarize_document(
        uploaded.read(), uploaded.filename, style, describe_images, socratic_tutor
    )
    if not result["ok"]:
        return jsonify(result), 400

    summary = result["summary"]
    original_stem = Path(uploaded.filename).stem or "document"
    response = jsonify({
        "ok": True,
        "summary": summary,
        "download_name": f"{original_stem}-simplified.docx",
    })
    return response


@app.post("/tutor")
def tutor():
    uploaded = request.files.get("document")
    question = request.form.get("question", "")
    try:
        history = request.form.get("history", "[]")
        history = json.loads(history)
    except (TypeError, ValueError):
        history = []

    if uploaded is None or not uploaded.filename:
        return jsonify({"ok": False, "error": "Upload a document before starting the tutor."}), 400

    result = tutor_chat(uploaded.read(), uploaded.filename, question, history)
    return jsonify(result), (200 if result["ok"] else 400)


def _download_text(filename: str, summary: dict) -> BytesIO:
    """Build one readable file containing both the bullets and full rewrite."""
    lines = [
        f"Simplified version of {filename}",
        "",
        "SUMMARY",
        f"Main idea: {summary['one_sentence']}",
        "",
        "Detailed summary:",
        summary["detailed_summary"],
        "",
        "IMAGE DESCRIPTIONS",
        *([f"- {description}" for description in summary["image_descriptions"]] or ["- None requested or found"]),
        "",
        "Key points:",
        *[f"- {point}" for point in summary["key_points"]],
        "",
        "Important words:",
        *[f"- {item['word']}: {item['meaning']}" for item in summary["important_words"]],
        "",
        "Next steps:",
        *([f"- {step}" for step in summary["next_steps"]] or ["- None listed"]),
        "",
        "FULL SIMPLIFIED DOCUMENT",
        summary["simplified_document"],
        "",
    ]
    return BytesIO("\n".join(lines).encode("utf-8"))


def _download_docx(filename: str, summary: dict) -> BytesIO:
    """Build a real Word document containing the summary and rewrite."""
    document = Document()
    document.add_heading(f"Simplified version of {filename}", level=0)
    document.add_heading("Summary", level=1)
    document.add_paragraph(f"Main idea: {summary['one_sentence']}")
    document.add_heading("Detailed summary", level=2)
    document.add_paragraph(summary["detailed_summary"])

    for heading, items in (
        ("Image descriptions", summary["image_descriptions"]),
        ("Key points", summary["key_points"]),
        ("Important words", [f"{item['word']}: {item['meaning']}" for item in summary["important_words"]]),
        ("Next steps", summary["next_steps"] or ["None listed"]),
        ("Socratic tutor questions", [
            f"{item['question']}\nHint: {item['hint']}\nSkill: {item['skill']}"
            for item in summary["tutor_questions"]
        ]),
    ):
        document.add_heading(heading, level=2)
        for item in items or ["None requested or found"]:
            document.add_paragraph(item, style="List Bullet")

    document.add_heading("Full simplified document", level=1)
    for paragraph in summary["simplified_document"].splitlines():
        document.add_paragraph(paragraph)

    output = BytesIO()
    document.save(output)
    output.seek(0)
    return output


@app.post("/download")
def download():
    payload = request.get_json(silent=True) or {}
    filename = payload.get("filename", "document")
    summary = payload.get("summary")
    if not isinstance(summary, dict) or not summary.get("simplified_document"):
        return jsonify({"ok": False, "error": "Generate a document before downloading."}), 400

    file_data = _download_docx(filename, summary)
    stem = Path(filename).stem or "document"
    return send_file(
        file_data,
        mimetype="application/vnd.openxmlformats-officedocument.wordprocessingml.document",
        as_attachment=True,
        download_name=f"{stem}-simplified.docx",
    )


@app.get("/")
def index():
    return render_template("index.html")


if __name__ == "__main__":
    app.run(debug=True, port=5000)
