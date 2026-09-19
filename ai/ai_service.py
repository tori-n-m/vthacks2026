"""
ai_service.py
-------------
All Gemini API calls live here. The backend only needs to import and call:

    summarize_document(file_bytes, filename, style="general")
    summarize_text(text, style="general")
    list_styles()

Every summarize function NEVER raises an error. It always returns a dict:

    success -> {"ok": True,  "summary": {...}, "error": None}
    failure -> {"ok": False, "summary": None,  "error": "message for the student"}

    summary = {
        "one_sentence":    "The main idea in one short sentence.",
        "key_points":      ["...", "..."],
        "important_words": [{"word": "...", "meaning": "..."}],
        "next_steps":      ["...", "..."],
    }

Keep this shape stable so the front end never breaks.
"""

import io
import json
import os
from pathlib import Path
from typing import List

from dotenv import load_dotenv
from google import genai
from google.genai import types
from pydantic import BaseModel

# Works both when run as a script (python ai/test_ai.py) and when the
# backend imports it (from ai.ai_service import ...).
try:
    from .prompts import (
        DEFAULT_STYLE,
        STYLE_LABELS,
        get_system_instruction,
        get_user_prompt,
    )
except ImportError:
    from prompts import (
        DEFAULT_STYLE,
        STYLE_LABELS,
        get_system_instruction,
        get_user_prompt,
    )

# Load GEMINI_API_KEY from the .env file in the repo root
load_dotenv(Path(__file__).resolve().parent.parent / ".env")
load_dotenv()  # also checks the current folder, doesn't override

# Model name can be changed in .env without editing code
MODEL_NAME = os.getenv("GEMINI_MODEL", "gemini-2.5-flash")

SUPPORTED_TYPES = {".txt", ".md", ".pdf", ".docx"}
MAX_FILE_MB = 10
MAX_CHARS = 60_000  # keeps requests small and fast


# ---------- The shape Gemini must answer in ----------

class ImportantWord(BaseModel):
    word: str
    meaning: str


class DocumentSummary(BaseModel):
    one_sentence: str
    key_points: List[str]
    important_words: List[ImportantWord]
    next_steps: List[str]


# ---------- Small helpers ----------

_client = None


def _get_client():
    """Create the Gemini client once, the first time it's needed."""
    global _client
    if _client is None:
        api_key = os.getenv("GEMINI_API_KEY")
        if not api_key:
            raise RuntimeError("GEMINI_API_KEY is missing. Add it to your .env file.")
        _client = genai.Client(api_key=api_key)
    return _client


def _error(message: str) -> dict:
    return {"ok": False, "summary": None, "error": message}


def _read_docx(file_bytes: bytes) -> str:
    """Pull plain text (paragraphs and tables) out of a Word document."""
    from docx import Document  # imported here so .txt/.pdf work without it

    doc = Document(io.BytesIO(file_bytes))
    parts = [p.text for p in doc.paragraphs if p.text.strip()]
    for table in doc.tables:
        for row in table.rows:
            parts.append(" | ".join(cell.text.strip() for cell in row.cells))
    return "\n".join(parts)


def _summarize(document_content, style: str) -> dict:
    """Send the document to Gemini and return the standard result dict."""
    if style not in STYLE_LABELS:
        style = DEFAULT_STYLE

    response = _get_client().models.generate_content(
        model=MODEL_NAME,
        contents=[get_user_prompt(), document_content],
        config=types.GenerateContentConfig(
            system_instruction=get_system_instruction(style),
            temperature=0.3,  # low = more faithful to the document
            response_mime_type="application/json",
            response_schema=DocumentSummary,
        ),
    )

    parsed = response.parsed
    if parsed is None and response.text:
        parsed = DocumentSummary(**json.loads(response.text))
    if parsed is None:
        return _error("The AI couldn't summarize this document. Please try another one.")

    return {"ok": True, "summary": parsed.model_dump(), "error": None}


def _friendly_error(exc: Exception) -> dict:
    """Turn a scary exception into a message a student can understand."""
    print(f"[ai_service] error: {exc!r}")  # full details for the developers
    text = str(exc)
    if "GEMINI_API_KEY" in text:
        return _error("The AI isn't set up yet (missing API key).")
    if "429" in text or "RESOURCE_EXHAUSTED" in text:
        return _error("The AI is busy right now. Please wait a minute and try again.")
    return _error("Something went wrong while summarizing. Please try again.")


# ---------- Functions the backend calls ----------

def list_styles() -> list:
    """For the front end dropdown: [{"id": "dyslexia", "label": "..."}, ...]"""
    return [{"id": key, "label": label} for key, label in STYLE_LABELS.items()]


def summarize_text(text: str, style: str = DEFAULT_STYLE) -> dict:
    """Summarize text the student pasted in."""
    try:
        if not text or not text.strip():
            return _error("There's no text to summarize.")
        return _summarize(text[:MAX_CHARS], style)
    except Exception as exc:
        return _friendly_error(exc)


def summarize_document(file_bytes: bytes, filename: str, style: str = DEFAULT_STYLE) -> dict:
    """Summarize an uploaded file (.pdf, .docx, .txt, .md)."""
    try:
        ext = Path(filename or "").suffix.lower()
        if ext not in SUPPORTED_TYPES:
            return _error("Sorry, I can only read PDF, Word (.docx), .txt, and .md files.")
        if not file_bytes:
            return _error("That file looks empty.")
        if len(file_bytes) > MAX_FILE_MB * 1024 * 1024:
            return _error(f"That file is too big. Please use one under {MAX_FILE_MB} MB.")

        if ext == ".pdf":
            # Gemini reads PDFs directly, so no extra library needed
            content = types.Part.from_bytes(data=file_bytes, mime_type="application/pdf")
        else:
            if ext == ".docx":
                text = _read_docx(file_bytes)
            else:
                text = file_bytes.decode("utf-8", errors="ignore")
            if not text.strip():
                return _error("I couldn't find any text in that file.")
            content = text[:MAX_CHARS]

        return _summarize(content, style)
    except Exception as exc:
        return _friendly_error(exc)
